from __future__ import annotations

import io
import json
import statistics
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import PROJECT_ROOT, settings
from app.db.models import Violation

FONT_CANDIDATES = [
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    Path("/Library/Fonts/Arial Unicode.ttf"),
    Path("C:/Windows/Fonts/DejaVuSans.ttf"),
    PROJECT_ROOT / "frontend" / "public" / "fonts" / "DejaVuSans.ttf",
]

VIOLATION_LABELS = {"weapon": "Оружие", "suspicious_bag": "Бесхозная сумка", "dog": "Собака"}


def _font_paths() -> tuple[Path | None, Path | None]:
    for path in FONT_CANDIDATES:
        if path.is_file():
            bold = path.with_name("DejaVuSans-Bold.ttf")
            return path, bold if bold.is_file() else path
    return None, None


FONT_REGULAR, FONT_BOLD = _font_paths()
FONT_NAME = "DejaVuSans" if FONT_REGULAR else "Helvetica"
FONT_BOLD_NAME = "DejaVuSans-Bold" if FONT_BOLD else "Helvetica-Bold"

if FONT_REGULAR:
    try:
        pdfmetrics.registerFont(TTFont(FONT_NAME, str(FONT_REGULAR)))
        pdfmetrics.registerFont(TTFont(FONT_BOLD_NAME, str(FONT_BOLD or FONT_REGULAR)))
    except Exception:  # pragma: no cover
        FONT_NAME = "Helvetica"
        FONT_BOLD_NAME = "Helvetica-Bold"


def _aware(value: datetime | None) -> datetime:
    if value is None:
        return datetime.now(timezone.utc)
    return value if value.tzinfo else value.replace(tzinfo=timezone.utc)


def watermark_text(user: str) -> str:
    return f"Лицей №6 «Перспектива» · {settings.PROJECT_NAME} · {user} · {datetime.now().strftime('%d.%m.%Y %H:%M')}"


def summary_for_rows(rows: list[Violation]) -> dict[str, Any]:
    by_type = {"weapon": 0, "suspicious_bag": 0, "dog": 0}
    acked = 0
    false_alarms = 0
    latencies: list[float] = []
    by_camera: dict[str, int] = {}
    for row in rows:
        by_type[row.violation_type] = by_type.get(row.violation_type, 0) + 1
        if row.is_acknowledged:
            acked += 1
        if row.is_false_alarm:
            false_alarms += 1
        if row.acknowledged_latency_sec is not None:
            latencies.append(float(row.acknowledged_latency_sec))
        name = row.camera.name if row.camera else f"Камера #{row.camera_id}"
        by_camera[name] = by_camera.get(name, 0) + 1

    overdue = sum(1 for value in latencies if value > settings.SLA_TARGET_SEC)
    return {
        "total": len(rows),
        "by_type": by_type,
        "acknowledged": acked,
        "false_alarms": false_alarms,
        "latency_avg": round(statistics.mean(latencies), 1) if latencies else None,
        "latency_median": round(statistics.median(latencies), 1) if latencies else None,
        "sla_target_sec": settings.SLA_TARGET_SEC,
        "sla_overdue": overdue,
        "top_cameras": sorted(by_camera.items(), key=lambda item: item[1], reverse=True)[:5],
    }


# ─── PDF ────────────────────────────────────────────────────────────────────

def _draw_watermark(canvas, doc, text: str) -> None:
    canvas.saveState()
    canvas.setFont(FONT_NAME, 22)
    canvas.setFillColor(colors.Color(0.35, 0.4, 0.5, alpha=0.13))
    canvas.translate(A4[0] / 2, A4[1] / 2)
    canvas.rotate(28)
    for offset in (-140, 40, 220):
        canvas.drawCentredString(0, offset, text)
    canvas.restoreState()
    canvas.setFont(FONT_NAME, 7)
    canvas.setFillColor(colors.grey)
    canvas.drawRightString(A4[0] - 12 * mm, 8 * mm, f"стр. {canvas.getPageNumber()}")


def build_pdf(rows: list[Violation], user: str, period_label: str, options: dict | None = None) -> bytes:
    options = options or {}
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title="Отчёт по инцидентам",
        author=user,
        subject=period_label,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleRu", parent=styles["Title"], fontName=FONT_BOLD_NAME, fontSize=16, leading=20)
    h2 = ParagraphStyle("H2Ru", parent=styles["Heading2"], fontName=FONT_BOLD_NAME, fontSize=11, leading=14, spaceBefore=6)
    body = ParagraphStyle("BodyRu", parent=styles["BodyText"], fontName=FONT_NAME, fontSize=8.5, leading=11)
    small = ParagraphStyle("SmallRu", parent=body, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#475569"))

    stamp = watermark_text(user)
    summary = summary_for_rows(rows)
    story: list[Any] = [
        Paragraph("Отчёт по инцидентам школьной системы безопасности", title_style),
        Paragraph(f"Лицей №6 «Перспектива» · период: {period_label}", body),
        Paragraph(f"Сформировал: {user} · {datetime.now().strftime('%d.%m.%Y %H:%M')}", small),
        Spacer(1, 5 * mm),
        Paragraph("Сводка", h2),
    ]

    sla_avg = f"{summary['latency_avg']:.0f} с" if summary["latency_avg"] is not None else "—"
    sla_med = f"{summary['latency_median']:.0f} с" if summary["latency_median"] is not None else "—"
    cards = [
        ["Всего инцидентов", "Оружие", "Бесхозные сумки", "Собаки", "Ложные"],
        [
            str(summary["total"]),
            str(summary["by_type"].get("weapon", 0)),
            str(summary["by_type"].get("suspicious_bag", 0)),
            str(summary["by_type"].get("dog", 0)),
            str(summary["false_alarms"]),
        ],
        ["Подтверждено", "Среднее время реакции", "Медиана", "Просрочено SLA", f"SLA-цель, с"],
        [str(summary["acknowledged"]), sla_avg, sla_med, str(summary["sla_overdue"]), str(summary["sla_target_sec"])],
    ]
    summary_table = Table(cards, colWidths=[42 * mm] * 5)
    summary_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), FONT_NAME),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#e2e8f0")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#0f172a")),
                ("FONTNAME", (0, 1), (-1, 1), FONT_BOLD_NAME),
                ("FONTNAME", (0, 3), (-1, 3), FONT_BOLD_NAME),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story += [summary_table, Spacer(1, 4 * mm), Paragraph("Инциденты", h2)]

    header = ["#", "Дата и время", "Камера", "Тип", "Уверенность", "Зона", "Статус", "Реакция, с"]
    data = [header]
    for row in rows[:600]:
        detected = _aware(row.detected_at).astimezone()
        latency = f"{row.acknowledged_latency_sec:.0f}" if row.acknowledged_latency_sec is not None else "—"
        data.append(
            [
                str(row.id),
                detected.strftime("%d.%m.%Y %H:%M:%S"),
                (row.camera.name if row.camera else f"Камера #{row.camera_id}")[:38],
                VIOLATION_LABELS.get(row.violation_type, row.violation_type),
                f"{row.confidence * 100:.1f}%",
                (row.zone_name or "—")[:20],
                "ложное" if row.is_false_alarm else ("отработано" if row.is_acknowledged else "ожидает"),
                latency,
            ]
        )
    table = Table(data, repeatRows=1, colWidths=[10 * mm, 34 * mm, 60 * mm, 34 * mm, 22 * mm, 32 * mm, 24 * mm, 20 * mm])
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), FONT_NAME),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD_NAME),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(table)
    if len(rows) > 600:
        story.append(Paragraph(f"Показаны первые 600 из {len(rows)} записей. Полный перечень — в Excel-выгрузке.", small))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("Документ содержит служебную информацию ограниченного доступа. Копирование и передача третьим лицам запрещены.", small))

    doc.build(story, onFirstPage=lambda canvas, doc_: _draw_watermark(canvas, doc_, stamp), onLaterPages=lambda canvas, doc_: _draw_watermark(canvas, doc_, stamp))
    return buffer.getvalue()


# ─── Excel ──────────────────────────────────────────────────────────────────

def build_excel(rows: list[Violation], user: str, period_label: str) -> bytes:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Сводка"
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    border = Border(bottom=Side(style="thin", color="CBD5E1"))

    summary_ws["A1"] = "Отчёт по инцидентам · Аналитика безопасности школы"
    summary_ws["A1"].font = Font(bold=True, size=15, color="1E3A8A")
    summary_ws["A2"] = f"Лицей №6 «Перспектива» · период: {period_label}"
    summary_ws["A3"] = watermark_text(user)
    summary_ws["A3"].font = Font(size=8, color="64748B")
    for col in ("A", "B"):
        summary_ws.column_dimensions[col].width = 34

    summary = summary_for_rows(rows)
    pairs = [
        ("Всего инцидентов", summary["total"]),
        ("Оружие", summary["by_type"].get("weapon", 0)),
        ("Бесхозные сумки", summary["by_type"].get("suspicious_bag", 0)),
        ("Собаки на территории", summary["by_type"].get("dog", 0)),
        ("Подтверждено", summary["acknowledged"]),
        ("Ложные срабатывания", summary["false_alarms"]),
        ("Среднее время реакции, с", summary["latency_avg"] if summary["latency_avg"] is not None else "—"),
        ("Медиана реакции, с", summary["latency_median"] if summary["latency_median"] is not None else "—"),
        ("Просрочено SLA", summary["sla_overdue"]),
        ("SLA-цель, с", summary["sla_target_sec"]),
    ]
    start = 5
    summary_ws.cell(row=start, column=1, value="Показатель")
    summary_ws.cell(row=start, column=2, value="Значение")
    for cell in summary_ws[start]:
        cell.fill = header_fill
        cell.font = header_font
    for index, (label, value) in enumerate(pairs, start=start + 1):
        summary_ws.cell(row=index, column=1, value=label)
        summary_ws.cell(row=index, column=2, value=value)
        for cell in summary_ws[index]:
            cell.border = border

    top_row = start + len(pairs) + 3
    summary_ws.cell(row=top_row, column=1, value="Топ камер по числу инцидентов").font = Font(bold=True)
    for index, (name, count) in enumerate(summary["top_cameras"], start=top_row + 1):
        summary_ws.cell(row=index, column=1, value=name)
        summary_ws.cell(row=index, column=2, value=count)

    ws = wb.create_sheet("Инциденты")
    columns = [
        ("ID", 8),
        ("Дата и время", 20),
        ("Камера", 32),
        ("Локация", 30),
        ("Тип", 20),
        ("Уверенность", 12),
        ("Трек", 14),
        ("Выдержка, с", 12),
        ("Зона", 22),
        ("A/B", 8),
        ("Статус", 14),
        ("Ложное", 10),
        ("Подтвердил", 18),
        ("Время реакции, с", 16),
    ]
    ws["A1"] = watermark_text(user)
    ws["A1"].font = Font(size=8, color="64748B")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    for index, (name, width) in enumerate(columns, start=1):
        cell = ws.cell(row=2, column=index, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}2"

    for row_index, row in enumerate(rows, start=3):
        detected = _aware(row.detected_at).astimezone()
        values = [
            row.id,
            detected.strftime("%d.%m.%Y %H:%M:%S"),
            row.camera.name if row.camera else f"Камера #{row.camera_id}",
            row.camera.location if row.camera else "",
            VIOLATION_LABELS.get(row.violation_type, row.violation_type),
            round(row.confidence * 100, 1),
            row.track_id or "",
            row.dwell_seconds if row.dwell_seconds is not None else "",
            row.zone_name or "",
            row.model_variant or "",
            "ложное" if row.is_false_alarm else ("отработано" if row.is_acknowledged else "ожидает"),
            "да" if row.is_false_alarm else "нет",
            row.acknowledged_by or "",
            round(row.acknowledged_latency_sec, 1) if row.acknowledged_latency_sec is not None else "",
        ]
        for col_index, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=col_index, value=value)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _row_dict(row: Violation) -> dict:
    return {
        "id": row.id,
        "detected_at": _aware(row.detected_at).isoformat(),
        "camera": row.camera.name if row.camera else f"Камера #{row.camera_id}",
        "violation_type": row.violation_type,
        "confidence": round(row.confidence, 4),
        "is_acknowledged": row.is_acknowledged,
        "is_false_alarm": row.is_false_alarm,
        "acknowledged_latency_sec": row.acknowledged_latency_sec,
        "zone_name": row.zone_name,
        "model_variant": row.model_variant,
    }


def rows_as_json(rows: list[Violation]) -> str:
    return json.dumps([_row_dict(row) for row in rows], ensure_ascii=False, default=str)
